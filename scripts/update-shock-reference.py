"""Derive two explicitly defined price-episode references; never change game parameters.

Requires openpyxl. --source-file supports a locally inspected annual Pink Sheet.
Thresholds and episode definition are declared before reading observations.
"""
import argparse
import hashlib
import io
import json
import math
from pathlib import Path
import urllib.request
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
URL = 'https://thedocs.worldbank.org/en/doc/74e8be41ceb20fa0da750cda2f6b9e4e-0050012026/related/CMO-Historical-Data-Annual.xlsx'
FIRST, LAST, THRESHOLD = 2000, 2025, 20


def price_episodes(values, first=FIRST, last=LAST, threshold=THRESHOLD):
    # Two preceding observations detect a run already in progress at firstYear.
    required = set(range(first - 2, last + 1))
    if not required <= values.keys() or any(not math.isfinite(values[y]) or values[y] <= 0 for y in required):
        raise ValueError('Missing, nonfinite or nonpositive annual index in reference window')
    change = {y: (values[y] / values[y - 1] - 1) * 100 for y in range(first - 1, last + 1)}
    return [{'year': y, 'increasePercent': change[y], 'aboveThreshold': change[y] >= threshold,
             'episodeStart': change[y] >= threshold and change[y - 1] < threshold}
            for y in range(first, last + 1)]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-file', type=Path)
    args = parser.parse_args()
    raw = args.source_file.read_bytes() if args.source_file else urllib.request.urlopen(URL, timeout=45).read()
    workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    sheet = workbook['Annual Indices (Nominal)']
    if sheet.cell(6, 3).value != 'Energy' or sheet.cell(8, 7).value != 'Food':
        raise ValueError('Unexpected index names/columns; review the source before updating')
    indices = {'energy': {}, 'food': {}}
    for row in sheet.iter_rows(min_row=10, values_only=True):
        if type(row[0]) is not int or not FIRST - 2 <= row[0] <= LAST:
            continue
        for kind, column in [('energy', 2), ('food', 6)]:
            if row[0] in indices[kind] or type(row[column]) not in (int, float):
                raise ValueError('Duplicate year or invalid index value')
            indices[kind][row[0]] = row[column]
    annual = {kind: price_episodes(values) for kind, values in indices.items()}
    means = {kind: sum(r['episodeStart'] for r in rows) / (LAST - FIRST + 1) for kind, rows in annual.items()}
    result = {'accessedOn': '2026-10-06', 'window': {'firstYear': FIRST, 'lastYear': LAST, 'years': LAST - FIRST + 1},
        'source': {'url': URL, 'landingPage': 'https://www.worldbank.org/en/research/commodity-markets',
            'workbookSha256': hashlib.sha256(raw).hexdigest(), 'workbookUpdatedLabel': sheet.cell(4, 1).value,
            'sheet': sheet.title, 'citation': 'World Bank Prospects Group, Commodity Price Data (The Pink Sheet), annual nominal energy and food indices.',
            'license': 'CC BY 4.0', 'licenseMetadata': 'https://datacatalog.worldbank.org/search/dataset/0038238/commodity-prices-history-and-projections'},
        'definition': {'thresholdIncreasePercent': THRESHOLD, 'direction': 'positive increases only',
            'episode': 'First year at or above threshold following a year below threshold. Consecutive qualifying years are one episode; carry-in from 1999 is excluded.',
            'scope': 'Price-rise proxy defined by MANDATO, not an official catalog of supply crises. Completed annual observations only; 2026 and forecasts excluded.'},
        'annual': annual, 'annualEpisodeMeans': means,
        'unmappedTypes': ['finance', 'interest-rates', 'pandemic', 'natural-disaster', 'semiconductor', 'migration'],
        'mappingLimit': 'A game shock has no observed price threshold and can represent supply disruption; its intensity is an index. These two price proxies diagnose frequency, not historical equivalence or all eight types. No parameter is fitted by this updater.'}
    (ROOT / 'docs/independent-shock-reference.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'updatedLabel': result['source']['workbookUpdatedLabel'], 'episodeMeans': means, 'hash': result['source']['workbookSha256']}))


if __name__ == '__main__':
    main()
